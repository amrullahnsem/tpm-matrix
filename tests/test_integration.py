"""Integration tests: repository, API and report generation.

Every test runs against a throwaway copy of the real database so the live
EFLOW data is never mutated by the test suite.
"""

from __future__ import annotations

import shutil
import sqlite3
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app import reports as report_builder
from app.repository import DB_PATH, Repository
from app.scoring import GapStatus, Priority

CRUSHER = "Crusher Maintenance"


@pytest.fixture()
def repo(tmp_path: Path) -> Repository:
    target = tmp_path / "matrix.db"
    shutil.copy(DB_PATH, target)
    return Repository(target)


@pytest.fixture()
def client(repo: Repository, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    import app.server as server

    monkeypatch.setattr(server, "repo", repo)
    monkeypatch.setattr(server, "OUTPUT_DIR", repo.db_path.parent / "reports")
    return TestClient(server.app)


class TestSchemaAndSeed:
    def test_expected_tables_exist(self, repo: Repository):
        with repo.connect() as conn:
            names = {
                r[0]
                for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")
            }
        assert {
            "positions", "competencies", "employees", "position_requirements",
            "assessments", "assessment_scores", "employee_current_scores", "training_records",
        } <= names

    def test_legacy_competencies_retained(self, repo: Repository):
        codes = {c.code for c in repo.list_competencies()}
        assert "TPM-AM-01" in codes, "the original 10 generic competencies must survive the migration"

    def test_thirteen_crusher_competencies_added(self, repo: Repository):
        crusher = [c for c in repo.list_competencies() if c.code.startswith("TPM-CR-")]
        assert len(crusher) == 13
        names = {c.name for c in crusher}
        assert "Crusher Shutdown SOP & Lock-Out Tag-Out (LOTO)" in names
        assert "Electrical MCC & Motor Condition Monitoring" in names

    def test_years_of_experience_populated(self, repo: Repository):
        for employee in repo.list_employees():
            assert employee.years_of_experience is not None
            assert employee.years_of_experience >= 0

    def test_axis_order_is_deterministic(self, repo: Repository):
        orders = [c.sort_order for c in repo.list_competencies()]
        assert orders == sorted(orders)
        assert len(set(orders)) == len(orders)

    def test_every_position_requirement_is_in_range(self, repo: Repository):
        with repo.connect() as conn:
            bad = conn.execute(
                "SELECT COUNT(*) FROM position_requirements"
                " WHERE required_rating < 0 OR required_rating > 5"
            ).fetchone()[0]
        assert bad == 0


class TestMatrixAndScoring:
    def test_matrix_axis_count(self, repo: Repository):
        scores, _, _ = repo.build_matrix("emp-101")
        assert len(scores) == 23

    def test_unassessed_competencies_are_none_not_zero(self, repo: Repository):
        scores, _, _ = repo.build_matrix("emp-101")
        by_id = {s.competency_id: s for s in scores}
        assert by_id["comp-101"].rating_score is None
        assert by_id["comp-1"].rating_score is not None

    def test_overall_matches_stored_baseline(self, repo: Repository):
        """emp-101's stored 78% must be reproducible from the score rows."""
        _, overall, _ = repo.build_matrix("emp-101")
        assert overall.percentage == pytest.approx(78.0, abs=0.01)
        assert overall.total_rating == 39

    def test_coverage_reflects_new_crusher_competencies(self, repo: Repository):
        _, overall, _ = repo.build_matrix("emp-101")
        assert overall.assessed_count == 10
        assert overall.unassessed_count == 13
        assert overall.coverage_percentage == pytest.approx(43.48, abs=0.01)

    def test_every_employee_builds_a_matrix(self, repo: Repository):
        for employee in repo.list_employees():
            scores, overall, gaps = repo.build_matrix(employee.id)
            assert scores and gaps
            assert 0 <= overall.percentage <= 100

    def test_position_specific_requirements_differ(self, repo: Repository):
        """Electrical specialist and mechanical technician need opposite skills."""
        _, _, gaps_elec = repo.build_matrix("emp-104")   # pos-4 Electrical
        _, _, gaps_mech = repo.build_matrix("emp-101")   # pos-1 Senior Mechanical
        elec = {g.competency_id: g.required_score for g in gaps_elec}
        mech = {g.competency_id: g.required_score for g in gaps_mech}
        assert elec["comp-113"] == 5 and mech["comp-113"] == 3
        assert elec["comp-112"] == 1 and mech["comp-112"] == 5

    def test_gap_results_are_priority_ordered(self, repo: Repository):
        _, _, gaps = repo.build_matrix("emp-101")
        ranks = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}
        seen = [ranks[str(g.priority)] for g in gaps]
        assert seen == sorted(seen)

    def test_unknown_employee_raises(self, repo: Repository):
        with pytest.raises(LookupError):
            repo.build_matrix("emp-does-not-exist")


class TestRecordAssessment:
    def test_writes_assessment_scores_and_current_scores(self, repo: Repository):
        before = repo.get_profile("emp-105")
        result = repo.record_assessment(
            employee_id="emp-105",
            assessor_name="Abdul Aziz Bin Abdullah (aziz)",
            cycle_name="2026 Q3 Periodic Review",
            assessment_date="2026-09-01",
            ratings={"comp-101": 5, "comp-102": 4, "comp-103": 3, "comp-104": 4},
        )
        assert result.total_rating == 16
        assert result.max_possible_rating == 20
        assert result.overall_percentage == 80.0

        with repo.connect() as conn:
            details = conn.execute(
                "SELECT COUNT(*) FROM assessment_scores WHERE assessment_id = ?", (result.id,)
            ).fetchone()[0]
        assert details == 4

        after = repo.get_profile("emp-105")
        assert after.overall.assessed_count == before.overall.assessed_count + 4
        by_id = {s.competency_id: s for s in after.scores}
        assert by_id["comp-101"].rating_score == 5

    def test_assessment_appears_in_history(self, repo: Repository):
        repo.record_assessment(
            employee_id="emp-102", assessor_name="Aziz", cycle_name="Cycle X",
            assessment_date="2026-09-02", ratings={"comp-101": 4},
        )
        history = repo.progress_history("emp-102")
        assert history[-1]["cycleName"] == "Cycle X"
        assert history[-1]["hasPerCompetencyDetail"] is True

    def test_resubmitting_updates_rather_than_duplicates(self, repo: Repository):
        repo.record_assessment(
            employee_id="emp-102", assessor_name="Aziz", cycle_name="A",
            assessment_date="2026-09-02", ratings={"comp-101": 2},
        )
        repo.record_assessment(
            employee_id="emp-102", assessor_name="Aziz", cycle_name="B",
            assessment_date="2026-09-03", ratings={"comp-101": 5},
        )
        with repo.connect() as conn:
            rows = conn.execute(
                "SELECT rating FROM employee_current_scores"
                " WHERE employee_id='emp-102' AND competency_id='comp-101'"
            ).fetchall()
        assert [r[0] for r in rows] == [5]

    def test_rejects_empty_ratings(self, repo: Repository):
        with pytest.raises(ValueError):
            repo.record_assessment(
                employee_id="emp-102", assessor_name="A", cycle_name="C",
                assessment_date="2026-01-01", ratings={},
            )

    def test_rejects_out_of_range_rating(self, repo: Repository):
        with pytest.raises(ValueError):
            repo.record_assessment(
                employee_id="emp-102", assessor_name="A", cycle_name="C",
                assessment_date="2026-01-01", ratings={"comp-101": 9},
            )

    def test_rejects_unknown_competency(self, repo: Repository):
        with pytest.raises(ValueError):
            repo.record_assessment(
                employee_id="emp-102", assessor_name="A", cycle_name="C",
                assessment_date="2026-01-01", ratings={"comp-nope": 3},
            )

    def test_rejects_unknown_employee(self, repo: Repository):
        with pytest.raises(LookupError):
            repo.record_assessment(
                employee_id="emp-nope", assessor_name="A", cycle_name="C",
                assessment_date="2026-01-01", ratings={"comp-1": 3},
            )

    def test_assessing_all_gaps_clears_action_items(self, repo: Repository):
        scores, _, gaps = repo.build_matrix("emp-105")
        requirements = {g.competency_id: g.required_score for g in gaps if g.required_score > 0}
        repo.record_assessment(
            employee_id="emp-105", assessor_name="Aziz", cycle_name="Full",
            assessment_date="2026-09-04", ratings=dict(requirements),
        )
        _, _, after = repo.build_matrix("emp-105")
        remaining = [g for g in after if g.is_actionable and g.required_score > 0]
        assert remaining == []
        assert all(g.status is GapStatus.ACHIEVED for g in after if g.required_score > 0)


class TestTrainingNeeds:
    def test_only_actionable_rows_returned(self, repo: Repository):
        rows = repo.training_needs()
        assert rows
        assert all(r["status"] in ("Gap", "Unassessed") for r in rows)

    def test_sorted_critical_first(self, repo: Repository):
        rows = repo.training_needs(department=CRUSHER)
        ranks = {str(Priority.CRITICAL): 0, str(Priority.HIGH): 1, str(Priority.MEDIUM): 2}
        seen = [ranks[str(r["priority"])] for r in rows]
        assert seen == sorted(seen)

    def test_department_filter_is_respected(self, repo: Repository):
        rows = repo.training_needs(department=CRUSHER)
        assert {r["department"] for r in rows} == {CRUSHER}

    def test_department_summary_totals(self, repo: Repository):
        summary = repo.department_summary(CRUSHER)
        assert summary.employee_count == 4  # emp-103 sits in Operations & TPM
        assert 0 < summary.average_percentage <= 100


class TestReports:
    def test_all_five_reports_generate_pdf_and_excel(self, repo: Repository, tmp_path: Path):
        artifacts = [
            report_builder.individual_competency_report(repo, "emp-101", tmp_path),
            report_builder.department_competency_report(repo, CRUSHER, tmp_path),
            report_builder.competency_gap_report(repo, tmp_path, department=CRUSHER),
            report_builder.training_needs_report(repo, tmp_path, department=CRUSHER),
            report_builder.competency_progress_report(repo, "emp-101", tmp_path),
        ]
        for artifact in artifacts:
            assert artifact.pdf and artifact.pdf.exists() and artifact.pdf.stat().st_size > 1000
            assert artifact.xlsx and artifact.xlsx.exists() and artifact.xlsx.stat().st_size > 1000
            assert artifact.pdf.read_bytes().startswith(b"%PDF")
            assert artifact.xlsx.read_bytes()[:2] == b"PK"

    def test_individual_report_includes_the_spider_chart(self, repo: Repository, tmp_path: Path):
        from reportlab.graphics import renderPDF
        from svglib.svglib import svg2rlg

        artifact = report_builder.individual_competency_report(repo, "emp-101", tmp_path)
        assert artifact.pdf.stat().st_size > 20000, "chart-less report is far smaller"

        profile = repo.get_profile("emp-101")
        assert renderPDF.drawToString(
            svg2rlg(__import__("io").BytesIO(
                __import__("app.radar", fromlist=["render_radar_svg"]).render_radar_svg(
                    profile.scores, size=1000, employee_name=profile.employee.full_name
                ).encode()
            ))
        )

    def test_excel_sheets_have_content(self, repo: Repository, tmp_path: Path):
        from openpyxl import load_workbook

        artifact = report_builder.individual_competency_report(repo, "emp-101", tmp_path)
        workbook = load_workbook(artifact.xlsx)
        assert {"Summary", "Competency Matrix", "Assessment History", "Training History"} <= set(
            workbook.sheetnames
        )
        assert workbook["Competency Matrix"].max_row == 24  # header + 23 competencies

    def test_training_needs_groups_by_competency(self, repo: Repository, tmp_path: Path):
        artifact = report_builder.training_needs_report(repo, tmp_path, department=CRUSHER)
        assert artifact.row_count < 68, "grouping must collapse per-staff rows"
        assert artifact.row_count > 0

    def test_progress_report_flags_missing_detail(self, repo: Repository, tmp_path: Path):
        artifact = report_builder.competency_progress_report(repo, "emp-101", tmp_path)
        from openpyxl import load_workbook

        rows = list(load_workbook(artifact.xlsx)["Progress Trend"].iter_rows(min_row=2, values_only=True))
        # The two pre-existing assessments have no assessment_scores rows.
        assert [r[7] for r in rows] == ["No", "No"]


class TestApi:
    def test_employee_list(self, client: TestClient):
        assert client.get("/").status_code == 200
        assert len(client.get("/api/employees").json()["employees"]) == 5

    def test_dashboard_and_chart(self, client: TestClient):
        assert client.get("/employee/emp-101").status_code == 200
        chart = client.get("/employee/emp-101/chart.svg")
        assert chart.status_code == 200
        assert chart.headers["content-type"].startswith("image/svg+xml")
        assert "<svg" in chart.text

    def test_assessment_form_renders_every_competency(self, client: TestClient):
        page = client.get("/assess/emp-101").text
        assert page.count("data-comp=") >= 23
        assert "live-pct" in page

    def test_preview_endpoint_matches_engine(self, client: TestClient):
        response = client.post(
            "/api/scoring/preview",
            json={"ratings": {"comp-1": 5, "comp-2": 4, "comp-3": 3},
                  "required": {"comp-1": 5, "comp-2": 4, "comp-3": 3}},
        )
        data = response.json()
        assert data["overallPercentage"] == 80.0
        assert data["coveragePercentage"] == 100.0

    def test_preview_counts_unrated_competencies(self, client: TestClient):
        data = client.post(
            "/api/scoring/preview",
            json={"ratings": {"comp-1": 5},
                  "required": {"comp-1": 5, "comp-2": 4, "comp-3": 3}},
        ).json()
        assert data["unassessedCount"] == 2
        assert data["coveragePercentage"] == pytest.approx(33.33, abs=0.01)

    def test_preview_ignores_competencies_with_no_requirement(self, client: TestClient):
        data = client.post(
            "/api/scoring/preview",
            json={"ratings": {"comp-1": 4}, "required": {"comp-1": 5, "comp-2": 0}},
        ).json()
        assert data["maxPossibleRating"] == 5
        assert data["unassessedCount"] == 0

    def test_submit_assessment_persists(self, client: TestClient, repo: Repository):
        before = repo.get_profile("emp-105").overall.assessed_count
        response = client.post(
            "/api/assessments",
            json={
                "employeeId": "emp-105",
                "assessorName": "Aziz",
                "cycleName": "2026 Q3",
                "assessmentDate": "2026-09-10",
                "ratings": {"comp-101": 5, "comp-102": 4, "comp-103": 3},
            },
        )
        assert response.status_code == 200
        assert response.json()["assessment"]["overallPercentage"] == 80.0
        assert repo.get_profile("emp-105").overall.assessed_count == before + 3

    def test_submit_rejects_bad_rating(self, client: TestClient):
        response = client.post(
            "/api/assessments",
            json={"employeeId": "emp-105", "assessorName": "A", "cycleName": "C",
                  "ratings": {"comp-101": 42}},
        )
        assert response.status_code == 400

    def test_submit_rejects_missing_field(self, client: TestClient):
        response = client.post("/api/assessments", json={"employeeId": "emp-105"})
        assert response.status_code == 400

    def test_404_for_unknown_employee(self, client: TestClient):
        assert client.get("/api/employees/emp-nope").status_code == 404
        assert client.get("/employee/emp-nope").status_code == 404

    def test_rating_scale_endpoint(self, client: TestClient):
        scale = client.get("/api/scoring/rating-scale").json()["scale"]
        assert [(s["rating"], s["percentage"]) for s in scale] == [
            (5, 100), (4, 80), (3, 60), (2, 40), (1, 20), (0, 0)
        ]

    def test_departments_endpoint(self, client: TestClient):
        departments = client.get("/api/departments").json()["departments"]
        assert {d["department"] for d in departments} == {"Crusher Maintenance", "Operations & TPM"}

    def test_report_downloads(self, client: TestClient):
        cases = [
            ("/reports/individual/download?employee=emp-101&fmt=pdf", b"%PDF"),
            ("/reports/individual/download?employee=emp-101&fmt=xlsx", b"PK"),
            (f"/reports/department/download?department={CRUSHER.replace(' ', '%20')}&fmt=pdf", b"%PDF"),
            ("/reports/gap/download?fmt=xlsx", b"PK"),
            ("/reports/training/download?fmt=xlsx", b"PK"),
            ("/reports/progress/download?employee=emp-101&fmt=pdf", b"%PDF"),
        ]
        for url, magic in cases:
            response = client.get(url)
            assert response.status_code == 200, url
            assert response.content.startswith(magic), url

    def test_report_download_validation(self, client: TestClient):
        assert client.get("/reports/individual/download?fmt=pdf").status_code == 400
        assert client.get("/reports/bogus/download?fmt=pdf").status_code == 404
        assert client.get("/reports/gap/download?fmt=csv").status_code == 400


class TestApiCasing:
    """Regression guard: every JSON key must be camelCase, never snake_case.

    Hand-built dicts and asdict() output used to disagree, which silently
    breaks the browser client.
    """

    @pytest.mark.parametrize(
        "url",
        [
            "/api/employees",
            "/api/employees/emp-101",
            "/api/employees/emp-101/gaps",
            "/api/competencies",
            "/api/departments",
            "/api/scoring/rating-scale",
        ],
    )
    def test_no_snake_case_keys_anywhere(self, client: TestClient, url: str):
        def walk(node):
            if isinstance(node, dict):
                for key, value in node.items():
                    assert "_" not in key, f"{url} returned snake_case key {key!r}"
                    walk(value)
            elif isinstance(node, list):
                for item in node:
                    walk(item)

        walk(client.get(url).json())

    def test_assessment_response_uses_camel_case(self, client: TestClient):
        body = client.post(
            "/api/assessments",
            json={"employeeId": "emp-105", "assessorName": "Aziz", "cycleName": "C",
                  "ratings": {"comp-101": 5, "comp-102": 4, "comp-103": 3}},
        ).json()
        assert body["assessment"]["overallPercentage"] == 80.0
        assert body["assessment"]["maxPossibleRating"] == 15
        assert body["assessment"]["cycleName"] == "C"

    def test_profile_and_employee_list_agree_on_casing(self, client: TestClient):
        employee = client.get("/api/employees").json()["employees"][0]
        profile = client.get(f"/api/employees/{employee['id']}").json()
        assert set(employee) <= set(profile["employee"])
        assert employee["employeeNumber"] == profile["employee"]["employeeNumber"]
        assert profile["assessments"][0]["assessmentDate"]

